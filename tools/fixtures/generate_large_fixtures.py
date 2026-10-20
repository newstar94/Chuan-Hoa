"""Đợt 0 — sinh tài liệu thử lớn cho Chuẩn Hóa (không chứa dữ liệu thật).

Tạo hai tệp:
  L1 ≈ 70 trang, ~5.000 đoạn, 25 bảng (5 bảng > 100 hàng có ô gộp, hàng tiêu đề 2 tầng),
     10 ảnh inline, 3 section khổ ngang, header/footer có số trang, 1 trường TOC.
  L2 ≈ 150 trang, ~12.000 đoạn (cùng cấu trúc, nhân bản nội dung).
Mỗi tệp cài sẵn các "bẫy" nhận diện đã nêu trong prompt sửa đổi; có thêm một bản
đầu văn bản gõ Unicode tổ hợp (NFD).

Cách dùng:
    pip install python-docx
    python tools/fixtures/generate_large_fixtures.py --out tmp/fixtures
"""

from __future__ import annotations

import argparse
import io
import struct
import unicodedata
import zlib
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

TRAPS = [
    "Căn cứ Hiến pháp nước Cộng hòa xã hội chủ nghĩa Việt Nam;",
    "Nhà thầu cam kết tuân thủ pháp luật nước Cộng hòa xã hội chủ nghĩa Việt Nam trong suốt quá trình thi công.",
    "Công trình tri ân các anh hùng liệt sĩ đã hy sinh vì độc lập dân tộc và hạnh phúc của nhân dân.",
]

LEGAL_LIST = [
    ("1", "Luật Xây dựng số 50/2014/QH13 do Quốc hội nước Cộng hòa xã hội chủ nghĩa Việt Nam ban hành"),
    ("2", "Luật Đấu thầu số 22/2023/QH15"),
    ("3", "Nghị định số 06/2021/NĐ-CP về quản lý chất lượng công trình"),
    ("4", "TCVN 4085:2025 Kết cấu gạch đá - Quy phạm thi công và nghiệm thu"),
]


def tiny_png() -> bytes:
    """A 2×2 grey PNG generated in memory (no binary asset in the repo)."""
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    raw = b"".join(b"\x00" + b"\x80\x80\x80" * 2 for _ in range(2))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0)) +
            chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText"); text.set(qn("xml:space"), "preserve"); text.text = instruction
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.append(begin); run._r.append(text); run._r.append(end)


def header_block(document: Document, nfd: bool) -> None:
    norm = (lambda s: unicodedata.normalize("NFD", s)) if nfd else (lambda s: s)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "CÔNG TY CỔ PHẦN XÂY DỰNG MẪU"
    table.cell(0, 1).text = norm("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    table.cell(1, 0).text = "Số: 01/TM-XD"
    table.cell(1, 1).text = norm("Độc lập - Tự do - Hạnh phúc")
    document.add_paragraph("Hưng Yên, ngày 20 tháng 10 năm 2026").alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = document.add_paragraph("BÁO CÁO")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].bold = True
    document.add_paragraph("Về thuyết minh biện pháp tổ chức thi công").alignment = WD_ALIGN_PARAGRAPH.CENTER


def data_table(document: Document, rows: int, index: int) -> None:
    table = document.add_table(rows=rows + 2, cols=5)
    table.style = "Table Grid"
    top = table.rows[0].cells
    top[0].merge(top[1]).text = f"Hạng mục {index}"
    top[2].merge(top[4]).text = "Kiểm tra chất lượng"
    sub = table.rows[1].cells
    for cell, label in zip(sub, ["STT", "Nội dung", "Biên bản", "Báo cáo", "Kế hoạch"]):
        cell.text = label
    for row in range(2, rows + 2):
        cells = table.rows[row].cells
        cells[0].text = str(row - 1)
        cells[1].text = f"Công tác kiểm tra số {row - 1} theo quy định hiện hành của hồ sơ mời thầu"
        cells[2].text = "Biên bản"
        cells[3].text = "Quy định"
        cells[4].text = "Số 1" if row == 2 else "Theo tiến độ"


def body(document: Document, count: int, offset: int) -> None:
    for i in range(count):
        n = offset + i
        if n % 400 == 0:
            document.add_paragraph(f"PHẦN {['I', 'II', 'III', 'IV', 'V'][(n // 400) % 5]}: THUYẾT MINH CHUNG")
        if n % 80 == 0:
            document.add_paragraph(f"{['I', 'II', 'III', 'IV', 'V', 'VI'][(n // 80) % 6]}. NỘI DUNG CHÍNH").runs[0].bold = True
        if n % 20 == 0:
            p = document.add_paragraph(f"{(n // 20) % 9 + 1}. Thông tin chung:")
            p.runs[0].bold = True
        if n % 10 == 0:
            document.add_paragraph(f"{(n // 20) % 9 + 1}.{(n // 10) % 2 + 1}. Yêu cầu kỹ thuật")
        text = ("- Nhà thầu tổ chức thi công theo đúng hồ sơ thiết kế, khoảng cách 10 m, "
                "giai đoạn 2020-2025, bảo đảm an toàn lao động và vệ sinh môi trường.")
        document.add_paragraph(text)
        if n % 250 == 0:
            for trap in TRAPS:
                document.add_paragraph(trap)


def build(path: Path, paragraphs: int, big_tables: int, small_tables: int, nfd: bool) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin, section.right_margin = Mm(30), Mm(15)
    section.top_margin = section.bottom_margin = Mm(20)
    add_field(section.header.paragraphs[0], "PAGE")
    add_field(section.footer.paragraphs[0], "PAGE")
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)

    header_block(document, nfd)
    add_field(document.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u')
    legal = document.add_table(rows=len(LEGAL_LIST) + 1, cols=2)
    legal.style = "Table Grid"
    legal.cell(0, 0).text, legal.cell(0, 1).text = "STT", "Văn bản pháp lý"
    for row, (number, name) in enumerate(LEGAL_LIST, start=1):
        legal.cell(row, 0).text, legal.cell(row, 1).text = number, name

    image = tiny_png()
    chunk = paragraphs // (big_tables + small_tables + 1)
    offset = 0
    landscape_left = 3
    for table_index in range(big_tables + small_tables):
        body(document, chunk, offset)
        offset += chunk
        if table_index < 10:
            document.add_picture(io.BytesIO(image), width=Mm(40))
            document.add_paragraph(f"Hình {table_index + 1}. Mặt bằng thi công")
        document.add_paragraph(f"Bảng số {table_index + 1:02d}: Kiểm soát chất lượng")
        if table_index % 8 == 7 and landscape_left > 0:
            new = document.add_section(WD_SECTION.NEW_PAGE)
            new.orientation = WD_ORIENT.LANDSCAPE
            new.page_width, new.page_height = Mm(297), Mm(210)
            landscape_left -= 1
        data_table(document, 110 if table_index < big_tables else 6, table_index + 1)
    body(document, chunk, offset)
    document.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", default="tmp/fixtures")
    arguments = parser.parse_args()
    out = Path(arguments.out)
    out.mkdir(parents=True, exist_ok=True)
    build(out / "L1_70trang.docx", 2300, 5, 20, nfd=False)
    build(out / "L1_70trang_NFD.docx", 2300, 5, 20, nfd=True)
    build(out / "L2_150trang.docx", 6000, 10, 40, nfd=False)
    print("Đã tạo:", ", ".join(sorted(p.name for p in out.glob("*.docx"))))


if __name__ == "__main__":
    main()
